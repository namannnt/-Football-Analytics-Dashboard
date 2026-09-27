from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from pptx import Presentation

import powerbi_dataset
from excel_export import build_excel
from export import build_pdf
from pptx_export import build_pptx


def sample_frame():
    return pd.DataFrame({"team": ["Atlas", "Boreal"], "points": [9, 6]})


def test_pdf_excel_and_pptx_smoke():
    frame = sample_frame()
    pdf = build_pdf("Demo", ["A deterministic takeaway."])
    excel = build_excel({"Standings": frame})
    pptx = build_pptx("Demo", ["A deterministic takeaway."])
    assert pdf.startswith(b"%PDF")
    assert load_workbook(BytesIO(excel))["Standings"].max_row == 3
    assert len(Presentation(BytesIO(pptx)).slides) == 2


def test_powerbi_csv_export(monkeypatch, tmp_path):
    monkeypatch.setattr(powerbi_dataset, "TABLES", ("weekly_standings", "missing"))
    monkeypatch.setattr(
        powerbi_dataset,
        "read_table_if_exists",
        lambda name: sample_frame() if name.endswith("weekly_standings") else pd.DataFrame(),
    )
    outputs = powerbi_dataset.export_powerbi_dataset(str(tmp_path), "csv")
    assert outputs == [tmp_path / "weekly_standings.csv"]
    assert pd.read_csv(outputs[0]).equals(sample_frame())
