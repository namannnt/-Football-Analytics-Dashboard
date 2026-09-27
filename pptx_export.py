"""PowerPoint summary export."""
from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.util import Pt


def build_pptx(title: str, takeaways: list[str]) -> bytes:
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = "Automated football performance brief"
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Key takeaways"
    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()
    for index, takeaway in enumerate(takeaways):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = takeaway
        paragraph.font.size = Pt(22)
        paragraph.space_after = Pt(12)
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()
